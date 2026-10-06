# -*- coding: utf-8 -*-
"""HTML rapordaki tablolari Excel sekmelerine cevirir (TR). Kaynak: uretilen HTML.
Sekme adi: <bolum no>.<tablo sirasi> <alt baslik>. Ilk sekme icindekiler. Notlar A2'de tek hucrede (Etiket: govde, etiket kalin coral)."""
import os, sys, re, math
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import veri
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
INK = "10332F"; BAS = "434343"; YES = "2E7D32"; KIR = "D32F2F"; COR = "FF7B52"; GOLD = "F5A623"; CT = "FFE3D8"; CD = "E85F36"; GRI = "F0EDE8"
CIZ = Border(top=Side(style="thin", color="E0E0E0"))
def F(b=False, c=INK, sz=10): return Font(name="Calibri", bold=b, color=c, size=sz)
def AL(h="left", w=True): return Alignment(horizontal=h, vertical="center", wrap_text=w)
AD = "VitrA_E-Ticaret_Buyume_Firsatlari"
H = open(os.path.join(veri.KOK, AD + ".html"), encoding="utf-8").read()
S = BeautifulSoup(H, "html.parser")
wb = Workbook(); wb.remove(wb.active)
ICS = wb.create_sheet("00 İçindekiler"); TOC = []
_num = re.compile(r"^[+\-]?%?\d[\d.]*(,\d+)?%?$")
_tarih = re.compile(r"^\d{1,2}\.\d{1,2}\.\d{2,4}$")
_km = re.compile(r"^([+\-]?\d[\d.]*(?:,\d+)?)\s?([KM])$")
def deger(t):
    """(deger, bicim) ; bicim: None | 'pct' | 'K' | 'M' ; tarih ve kod metin kalir."""
    t = " ".join(t.split())
    if not t or t == "-" or _tarih.match(t): return t, None
    if _num.match(t):
        s = t.replace("%", "").replace(".", "").replace(",", ".")
        try:
            v = float(s)
            return (int(v) if v.is_integer() and "," not in t else v), ("pct" if "%" in t else None)
        except ValueError: return t, None
    m = _km.match(t)
    if m:
        try: return float(m.group(1).replace(".", "").replace(",", ".")) * (1e3 if m.group(2) == "K" else 1e6), m.group(2)
        except ValueError: return t, None
    return t, None
def yaz(ws, r, c, v, bold=False, renk=INK, h="left", fill=None):
    cell = ws.cell(row=r, column=c, value=v); cell.font = F(bold, renk); cell.alignment = AL(h); cell.border = CIZ
    if fill: cell.fill = PatternFill("solid", fgColor=fill)
    return cell
_DUR = {"ve", "ile", "en", "ya", "da", "de", "için", "veya", "ki", "bir", "mi", "ne", "hangi", "the", "and", "of"}
def kisalt(t, n_):
    t = re.sub(r"[\[\]:*?/\\]", "", " ".join(t.split())).strip()
    if len(t) <= n_: return t
    w = t[:n_ + 1].split(" ")
    w = w[:-1] if len(w) > 1 else [t[:n_]]
    while len(w) > 1 and (w[-1].lower().strip(",·-&") in _DUR or not w[-1].strip(",·-&")): w.pop()
    return " ".join(w).rstrip(" -·,&")
_KOR = InlineFont(rFont="Calibri", sz=10, b=True, color=COR); _GOV = InlineFont(rFont="Calibri", sz=10, color=INK)
def notlar_yaz(ws, notlar, genislik_top):
    parca = []
    for i, nt in enumerate(notlar):
        et, _, gv = nt.partition(": ")
        if not gv: et, gv = "", nt
        if et: parca.append(TextBlock(_KOR, et + ": "))
        parca.append(TextBlock(_GOV, gv + ("\n" if i < len(notlar) - 1 else "")))
    c = ws.cell(row=2, column=1); c.value = CellRichText(parca); c.alignment = Alignment(vertical="top", wrap_text=True); c.border = CIZ
    satir = sum(max(1, math.ceil(len(nt) * 1.08 / max(40, genislik_top))) for nt in notlar)
    ws.row_dimensions[2].height = min(409, 14 * satir + 8)
def sekme(ad, baslik, notlar, basliklar, satirlar, genislik=None, renkler=None, metin_sutun=()):
    ws = wb.create_sheet(ad)
    yaz(ws, 1, 1, baslik, True); ws.row_dimensions[1].height = 22
    n_ = len(basliklar)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(n_, 2))
    hr = 4
    for i, b in enumerate(basliklar, 1):
        cc = ws.cell(row=hr, column=i, value=b); cc.font = F(True, "FFFFFF"); cc.fill = PatternFill("solid", fgColor=BAS); cc.alignment = AL("center"); cc.border = CIZ
    ws.row_dimensions[hr].height = 30
    w = genislik or [max(12, min(60, max([len(str(b))] + [len(str(s[i])) for s in satirlar if i < len(s)]) + 2)) for i, b in enumerate(basliklar)]
    for i, ww in enumerate(w, 1): ws.column_dimensions[get_column_letter(i)].width = ww
    notlar_yaz(ws, notlar, sum(w[:max(n_, 2)]) if len(w) >= 2 else 80)
    for j, sat in enumerate(satirlar):
        rr = hr + 1 + j; yuk = 15
        for i, dv in enumerate(sat, 1):
            metin_kal = (i - 1) in metin_sutun or (i <= n_ and re.search(r"^(Kod|SKU|Model kodu|Model|Kullanıldığı bölüm|Bölüm|Tarih)$", str(basliklar[i - 1])))
            v, bic = (dv, None) if (metin_kal or not isinstance(dv, str)) else deger(dv)
            sayisal = isinstance(v, (int, float))
            cell = yaz(ws, rr, i, v, h="center" if sayisal else "left")
            imli = isinstance(dv, str) and dv.strip()[:1] in "+-" and len(dv.strip()) > 1
            if bic == "pct":
                ond = "," in dv
                cell.number_format = ('+"%"0.0;-"%"0.0;"%"0' if ond else '+"%"0;-"%"0;"%"0') if imli else ('"%"0.0' if ond else '"%"0')
            elif bic in ("K", "M"):
                ond = "," in dv
                cell.number_format = ("#,##0.0" if ond else "#,##0") + (',"K"' if bic == "K" else ',,"M"')
            elif isinstance(v, int): cell.number_format = "#,##0"
            elif isinstance(v, float): cell.number_format = "#,##0.0"
            rk = renkler[j][i - 1] if renkler and j < len(renkler) and i - 1 < len(renkler[j]) else None
            if rk and sayisal and v != 0: cell.font = F(True, YES if rk == "up" else KIR)
            if isinstance(dv, str) and dv.startswith("Öncelik"):
                p = {"Öncelik 1": (GOLD, INK), "Öncelik 2": (CT, CD), "Öncelik 3": (GRI, "4A4A4A")}[dv]
                cell.fill = PatternFill("solid", fgColor=p[0]); cell.font = F(True, p[1])
            if isinstance(v, str) and v: yuk = max(yuk, 13 * math.ceil(len(v) / max(8, w[i - 1] - 2)) + 4)
        ws.row_dimensions[rr].height = min(yuk, 300)
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    if len(satirlar) > 12 and n_:
        ws.auto_filter.ref = "A%d:%s%d" % (hr, get_column_letter(n_), hr + len(satirlar))
    return ws
def metin(el):
    el = BeautifulSoup(str(el), "html.parser")
    for b in el.select("button"): b.decompose()   # acilir liste ve ok dugmeleri
    for s in el.select("span.t"): s.replace_with(s.get_text())   # dil katmani
    for sup in el.select("sup.ref"):
        sup.replace_with(" [Kaynakça %s]" % ", ".join(a.get_text() for a in sup.select("a")))
    for kl in el.select(".kwlist"):
        ogeler = []
        for kw in kl.select(".kw"):
            ic = kw.find("span")
            if ic is not None:
                v_ = " ".join(ic.get_text(" ").split()); ic.extract(); ogeler.append("%s (%s)" % (" ".join(kw.get_text(" ").split()), v_))
            else: ogeler.append(" ".join(kw.get_text(" ").split()))
        kl.replace_with(" · ".join(ogeler))
    t = " ".join(el.get_text(" ").split())
    t = re.sub(r"\s+([,.;:)])", r"\1", t); t = re.sub(r"\(\s+", "(", t)
    return t.replace("–", "-").replace("—", "-")
def hucre(td):
    dv = td.select_one("[data-v]")
    if dv is not None and td.get_text(" ", strip=True) == dv.get_text(" ", strip=True): return dv["data-v"]
    return metin(td)
def renk(td):
    if td.select_one("span.up"): return "up"
    if td.select_one("span.dn"): return "dn"
    return None
# --- bolumler -> sekmeler ---
AD_DUZ = {"02.3 Banyo yenileme · Karşılığı": "02.3 Banyo yenileme senaryosu", "02.4 Banyo yenileme · Değer": "02.4 Banyo yenileme girdileri",
          "13.2 Rakip · Set paket": "13.2 Rakip modeller matrisi", "13.3 Rakip · Set ve paket": "13.3 Rakip modeller detayı",
          "18.3 Çok satanlar · Kategori": "18.3 Çok satanlar · Adet", "18.4 Çok · Ciroya en çok": "18.4 Çok satanlar · Ciro payı", "18.5 Çok · Fiyat bandına": "18.5 Çok satanlar · Fiyat bandı",
          "18.10 Değerlendi · Yorum teması": "18.10 Değerlendirme · Yorum", "18.11 Değerlendir · Soru teması": "18.11 Değerlendirme · Soru", "18.12 Değerlendi · Yanıt kalıbı": "18.12 Değerlendirme · Yanıt",
          "18.13 Hepsibur · Trendyol resmi": "18.13 Hepsiburada · Kıyas", "18.14 Hepsiburada satış · Kod": "18.14 Hepsiburada · Ürün satış",
          "09.3 30 YouTube arama ifadesi": "09.3 YouTube ifade listesi",
          }
# numaradan bağımsız düzeltme: sekme adının "NN.k " sonrası
AD_SON = {"Marka · Üretici markalar": "Marka adı · Üretici", "Marka adı · Perakendeci": "Marka adı · Perakende",
          "Marka x · Aylık arama": "Marka x kat. · Hacim", "Marka x · YoY değişim": "Marka x kat. · YoY", "Marka x · 2023'ten bu yana": "Marka x kat. · 2023'ten",
          "YZ payına · Tüm yazılar": "YZ payı · Tümü", "YZ payına · Sıralaması": "YZ payı · Sıra korunan",
          "AI Overview · Arama grubu": "AIO · Arama grupları", "AI · Arama ifadesi": "AIO · Örnek aramalar",
          "GA4 aylık görünüm · Oturum": "GA4 aylık · Oturum", "GA4 · Gelire göre ilk": "GA4 giriş · İlk 15 gelir", "GA4 · Oturuma göre ilk": "GA4 giriş · İlk 15 oturum",
          "GA4 giriş · Giriş sayfası": "GA4 giriş · Sayfa türü", "GA4 giriş · Önceki sayfa": "GA4 giriş · Önceki sayfa", "GA4 site · Arama terimi": "GA4 arama · Terim", "GA4 site içi · Terim türü": "GA4 arama · Terim türü"}

_GENEL = {"Oyuncu", "Gösterge", "Ürün", "Kategori", "Tema", "Marka", "Alan adı", "Yıllık yenileme oranı"}
KISA = {"Yıllık banyo yenileme ve yeni konut banyosu tahmini": "Banyo yenileme tahmini", "Kullanıcı hangi özellikle arıyor? SSG": "Özellik araması SSG",
        "Kullanıcı hangi özellikle arıyor? BM": "Özellik araması BM", "İhtiyaç Dili: Kullanıcı Ne Arıyor?": "İhtiyaç sınıfları",
        "Sitenin fiilen aldığı trafik: Search Console sorguları": "Search Console sorguları", "Trafik hangi sayfa türlerine geliyor?": "Sayfa türleri",
        "Autocomplete önerilerinin tema oranı": "Autocomplete tema oranı", "En çok tık getiren sorgular · Search Console": "En çok tık getiren sorgular",
        "A · VitrA gamı: kategori bazında VitrA nerede?": "A · Kategori bazında VitrA", "A · VitrA'nın ilk 20'de görünmediği kelimeler": "A · İlk 20 dışı kelimeler",
        "B · Yakın kategori fırsatları: aramayı kim kazanıyor?": "B · Yakın kategoriler", "AI Overview'da en çok kaynak gösterilen alan adları": "AI Overview kaynakları",
        "YouTube: Montaj, Tamir ve Karar Videoları": "YouTube aramaları", "Genişletilmiş tarama: 68 arama, altı niyet grubu": "YouTube genişletilmiş tarama",
        "Katalog ve Talep Eşleşmesi": "Katalog ve talep", "Yeni Kategori ve Segment Fırsatları": "Yeni kategori temaları",
        "Kanal ölçeği: pazaryeri, yapı market, pure player ve marka siteleri": "Kanal ölçeği", "Kategori baş kelimeleri: hacim, potansiyel ve VitrA sırası": "Baş kelimeler ve VitrA sırası",
        "Rehber içerik: rakiplerin trafik aldığı konular": "Rehber içerik konuları", "Hedef kategoriler: fiyat bandı ve VitrA'nın yeri": "Hedef kategoriler",
        "Hedef dışı kesitler: hacim, marka yoğunluğu ve VitrA'ya uyum": "Hedef dışı kesitler", "Alt kesitler: vitra.com.tr, pazaryeri ve fiyat karşılaştırma siteleri": "Alt kesitler",
        "vitra.com.tr kategori sayfaları: ürün, fiyat aralığı ve stok": "vitra.com.tr stok", "Aynı model kodunda vitra.com.tr ve pazaryeri fiyatı": "Model kodu fiyat farkı",
        "Çok satan ürünler: adet ve ciro katkısı": "Çok satanlar", "Görüntülenmeden satışa: ürün dönüşümü": "Dönüşüm",
        "Değerlendirmeler ve sorular: memnuniyet ve şikayet temaları": "Değerlendirmeler", "Hepsiburada: son 12 ay satış ve Trendyol ile karşılaştırma": "Hepsiburada satış",
        "Trendyol'un Enleri: kategori listelerinde VitrA (Eylül 2026)": "Trendyol Enleri", "Aynı ürünün satıcıları: 24 VitrA ve Artema listesi": "Aynı ürünün satıcıları",
        "Trendyol ve Hepsiburada çok satan listeleri: 14 kategori": "Çok satan listeleri", "VitrA ürünlerinde kanal fiyat farkı: 29 ürün": "29 ürün fiyat farkı",
        "Akakçe ve Cimri: 29 VitrA ürününde satıcı sayısı ve en düşük fiyat": "Akakçe ve Cimri 29 ürün", "Fiyat karşılaştırma sitelerinde kategori görünümü": "Fiyat karşılaştırma kategorileri",
        "Site içi arama: kullanıcının yazdığı ifadeler ne döndürüyor?": "Site içi arama", "Sepet ve ödeme: üyeliksiz alışverişte ek adımlar": "Sepet ve ödeme",
        "Yolculuk adımları: gözlem, etki ve öneri": "Yolculuk adımları", "Hacimli ve ölçüye bağlı ürünlerde etkileşim modeli": "Hacimli ürünlerde etkileşim",
        "Ölçüm setleri: dört veri seti ve örneklemleri": "GEO ölçüm setleri",
        "Aylık organik performans: 2025 ve 2026": "Aylık organik 2025-2026", "Trafik hangi sayfa türlerine geliyor? · Oca-Eyl 2026": "Sayfa türleri 2026",
        "Kategori bazında organik click · Oca-Eyl 2026": "Kategori click 2026", "En çok click alan sayfalar · Oca-Eyl 2026": "En çok click alan sayfalar",
        "Blog (İlham Veren Fikirler): aylık click, konu grupları ve yazılar": "Blog",
        "Yapay zeka özelliklerinde gösterim: site geneli ve blog": "Yapay zeka gösterim payı", "Sayfa türüne göre yapay zeka gösterimi": "YZ sayfa türü", "Yapay zeka özelliklerinde en çok gösterilen sayfalar · site geneli": "YZ en çok gösterilen sayfa", "Yapay zeka gösterim payına göre blog yazılarında click değişimi": "YZ payına göre",
        "Yapay zeka özelliklerinde gösterim alan sayfalar ve click değişimi": "YZ gösterimi alan sayfalar", "AI Overview çıkan ve çıkmayan blog aramalarında click ve sıra": "AI Overview",
        "Aylık genel görünüm: gelir, satın alma ve oturum": "GA4 aylık görünüm", "Kanal performansı: oturum, satın alma ve gelir": "GA4 kanal",
        "Site içi arama terimleri": "GA4 site içi arama", "Giriş sayfası (landing page) türüne göre gelir": "GA4 giriş sayfası", "Ödeme ve teslimat adımı": "GA4 ödeme ve teslimat", "Satın alma dışı talep olayları": "GA4 talep olayları",
        "Blog ve koleksiyon sayfaları: görüntüleme": "GA4 blog ve koleksiyon", "Promosyon ve kupon": "GA4 promosyon",
        "Yapay zeka yanıt takibi · 111 markasız soruda markaların anılma payı": "AI yanıtlarında markalar",
        "Yapay zeka yanıt takibi · 111 markasız soruda kaynak gösterilen alan adları": "AI yanıt kaynakları",
        "Yapay zeka yanıt takibi · 29 satın alma ve montaj sorusunda VitrA ve vitra.com.tr": "Satın alma ve montaj soruları",
        "SEOmonitor takibi · 2.140 kelimede AI Overview ve VitrA": "SEOmonitor AI Overview",
        "Rapor hedef kelimeleri · 322 kelimede AI Overview ve kaynak siteler": "Hedef kelimelerde AIO",
        "Sitenin hazırlığı · Search Console'da rehber sayfaları ve 21 soru sorgusu": "GSC rehber ve soru sorguları",
        "Marka adıyla yapılan aramalar: üreticiler ve perakendeciler": "Marka adı",
        "Marka + kategori aramaları: markaların kategori bazında aranma hacmi": "Marka x kategori",
        "\"vitra\" + kategori aramaları: yıllar ve kategoriler": "vitra + kategori yıllar",
        "\"vitra\" ile birlikte aranan ifadeler: ihtiyaç sınıfları": "vitra ihtiyaç sınıfları",
        "Google'da click payı: vitra.com.tr ve rakipler": "Google click payı", "Marka siteleri arasında organik ziyaret": "Marka siteleri ziyaret",
        "Kart harcamaları: banyo ile ilişkili sektörler": "Kart harcamaları", "Konut hareketliliği ve kredi koşulları": "Konut ve kredi",
        "Soru sorguları: gösterim, tık ve sıra": "Soru sorguları", "Marka sayfası göstergeleri ve rakipler": "Şikayetvar marka göstergeleri",
        "Şikayet temaları ve örnek alıntılar": "Şikayet temaları", "Altı aylık dönemlerde tema geçişi": "Altı aylık tema geçişi",
        "Kanal ölçeği": "Kanal ölçeği", "Pazaryerindeki görünüm: Trendyol": "Trendyol görünümü", "Rakip ve benzer modeller": "Rakip modeller",
        "Banyo trafiği hangi sitelerde?": "Banyo trafiği", "Alt kategori trafiği: Trendyol ve Hepsiburada": "Alt kategori trafiği",
        "Kategori yapısı ve VitrA'nın listeleme payı": "Listeleme payı", "Satıcı yapısı, hizmet ve arama önerileri": "Satıcı yapısı",
        "Düşük fiyatlı, yüksek trafikli kategoriler": "Düşük fiyat yüksek trafik", "Ürün tipi ve mecra bazında fiyat bandı": "Mecra fiyat bandı",
        "VitrA ve Artema ilanlarında mecra farkı": "VitrA mecra farkı", "VitrA sitesinin göründüğü kelimeler": "Shopping VitrA kelimeleri",
        "Ürün grubuna göre fark medyanı": "Ürün grubu fark medyanı", "Kargo, teslimat, iade ve montaj": "Kargo ve iade",
        "Garanti, yedek parça ve servis iletişimi": "Garanti ve servis", "Google Business Profile: mağaza ve servis noktaları": "Business Profile",
        "Ürün keşfi ve dijital deneyim örnekleri": "Dijital deneyim örnekleri", "Duygu dağılımı: marka bazında": "Duygu dağılımı",
        "Pain point profili: kategori geneli mi, markaya özgü mü?": "Pain point profili", "Kategori bazında adet, ciro payı ve memnuniyet": "Kategori satış payı",
        "Çeyreklik satış, fiyat ve karışım": "Çeyreklik satış", "Müşteri profili ve sipariş yapısı": "Müşteri profili", "Rakip perakendeci ve marka mağazaları": "Rakip mağazalar",
        "Alt kesitlerde rakip mağazalar": "Alt kesit rakip mağazalar", "Aynı aramalarda görünen siteler": "Organik rakipler", "Site trafiği ve kanal kırılımı": "Similarweb trafik",
        "Tema bazında VitrA ve en yüksek rakip": "Tema bazında rakip", "Marka arama talebi": "Marka arama talebi", "Video fırsatları": "Video fırsatları"}
_RANK_BAS = {"serp": ("Alan adı", "Kaynak gösterildiği kelime"), "rakip": ("Alan adı", "Tahmini aylık organik ziyaret")}
def _h3_bolge(el, sec):
    h3 = el.find_previous("h3")
    return h3 if h3 is not None and h3.find_parent("section") is sec else None
def _kaynak(el, sec, h3):
    """Tablodan sonraki ilk kaynak satiri (ayni alt baslik icinde); yoksa bolumdeki sonraki, o da yoksa ilk kaynak."""
    for src in el.find_all_next("p", class_="src"):
        if src.find_parent("section") is not sec: break
        if _h3_bolge(src, sec) is h3: return metin(src)
        break
    for src in el.find_all_next("p", class_="src"):
        if src.find_parent("section") is sec: return metin(src)
        break
    ilk = sec.select_one("p.src"); return metin(ilk) if ilk else ""
def _ayirt(t, wrap):
    pan = t.find_parent("div", attrs={"role": "tabpanel"})
    if pan is not None and pan.get("id"):
        btn = S.find(attrs={"aria-controls": pan["id"]})
        if btn is not None: return " ".join(btn.get_text(" ").split())   # metin() düğmeleri sildiği için sekme adı doğrudan okunur
    tb = wrap.find_previous("p", class_="tbas")
    if tb is not None and _h3_bolge(tb, wrap.find_parent("section")) is _h3_bolge(wrap, wrap.find_parent("section")): return metin(tb)
    ths = [metin(x_) for x_ in t.select("thead th")]
    if ths and ths[0] in _GENEL and len(ths) > 1: return ths[1]
    return ths[0] if ths else ""
for sec in S.select("main section"):
    sid = sec.get("id"); h2 = sec.find("h2")
    no_ = metin(h2.select_one(".no")) if h2 and h2.select_one(".no") else ""
    bas = metin(h2).split(" ", 1)[1] if h2 else sid
    # gizli sekmedeki tablo aktarılmaz; "xl" sınıflı sarmalayıcıdaki tablolar (ör. marka aramalarının perakendeci ve değişim sekmeleri) her sekmeyle ayrı sayfaya yazılır
    gizli = lambda t_: t_.find_parent("div", attrs={"role": "tabpanel"}) is not None and t_.find_parent("div", attrs={"role": "tabpanel"}).has_attr("hidden") and "xl" not in (t_.find_parent("div").get("class") or [])
    ogeler = [t_ for t_ in sec.select("div.tw table, ol.rank")
              if not gizli(t_) and not (t_.find_parent("dialog") and sid in ("talep", "serp"))]
    if not ogeler: continue
    h3_say = {}
    for t_ in ogeler: h3_say[id(_h3_bolge(t_, sec))] = h3_say.get(id(_h3_bolge(t_, sec)), 0) + 1
    for ti, t in enumerate(ogeler, 1):
        dlg = t.find_parent("dialog")
        wrap = t if t.name == "ol" else t.find_parent("div")
        h3 = _h3_bolge(wrap, sec)
        _dh = dlg.select_one(".pophd b") or dlg.find(["h3", "h4", "strong"]) if dlg is not None else None
        alt = metin(_dh) if _dh is not None else (metin(h3) if h3 else "")
        alt = re.sub(r"\s*×\s*$", "", alt)
        renkler = []
        if t.name == "ol":
            b1, b2 = _RANK_BAS.get(sid, ("Ad", "Değer"))
            thead = ["Sıra", b1, b2]; aciklama = []
            rows = [[metin(li.select_one(".rn")), metin(li.select_one(".rl")), metin(li.select_one(".rv"))] for li in t.select("li")]
        else:
            thead = [metin(th) for th in t.select("thead th")]
            aciklama = ["%s: %s" % (metin(th) or "-", th.get("data-t", "")) for th in t.select("thead th") if th.get("data-t")]
            rows = []
            for tr in t.select("tbody tr"):
                tds = tr.select("td"); rows.append([hucre(td) for td in tds]); renkler.append([renk(td) for td in tds])
        h3n = h3.find_next_sibling("p", class_="h3n") if h3 is not None else None
        notlar = ["Bölüm: %s %s%s" % (no_, bas, (" · " + alt) if alt else "")]
        if h3n is not None: notlar.append("Açıklama: " + metin(h3n))
        else:
            lede = sec.select_one("p.lede")
            if lede: notlar.append("Kapsam: " + metin(lede))
        kay = _kaynak(wrap, sec, h3)
        if kay: notlar.append(kay if kay.startswith("Kaynak") else "Kaynak: " + kay)
        if aciklama: notlar.append("Sütun açıklamaları: " + " | ".join(aciklama))
        etiket = KISA.get(alt, alt) or KISA.get(bas, bas)
        on = "%s.%d " % (no_, ti)
        if h3_say.get(id(h3), 0) > 1 and dlg is None:
            ay = kisalt(_ayirt(t, wrap), 16) if t.name != "ol" else "Sıralama"
            ad = on + kisalt(etiket, max(8, 31 - len(on) - len(ay) - 3)) + " · " + ay
            baslik = "%s · %s · %s" % (bas, etiket, _ayirt(t, wrap) if t.name != "ol" else "Sıralama")
        else:
            ad = on + kisalt(etiket, 31 - len(on)); baslik = "%s · %s" % (bas, etiket)
        ad = kisalt(ad, 31)
        _on, _, _son = ad.partition(" ")
        ad = AD_DUZ.get(ad) or ((_on + " " + AD_SON[_son]) if _son in AD_SON else ad)
        sekme(ad, baslik, notlar, thead, rows, renkler=renkler)
        TOC.append((ad, "%s %s" % (no_, bas), alt or "-", len(rows)))
# kaynakca
rows = []; lis = S.select("ol.kaynakca li")
for li in lis:
    kb = li.select_one(".kb"); a = li.select("a.u")
    rows.append([metin(li.select_one(".kn")), metin(kb) if kb else "", " · ".join(x_.get("href") for x_ in a)])
ws = sekme("Kaynakça", "Kaynakça", ["Okuma: raporun metnindeki üst simge numaraları bu listeye bağlanmaktadır; Adres sütunundaki ilk bağlantı tıklanabilir."], ["No", "Kaynak", "Adres"], rows, [6, 90, 70], metin_sutun=(0,))
for j, li in enumerate(lis):
    a = li.select_one("a.u")
    if a is not None and a.get("href", "").startswith("http"):
        c_ = ws.cell(row=5 + j, column=3); c_.hyperlink = a["href"]; c_.font = Font(name="Calibri", size=10, color=INK, underline="single")
TOC.append(("Kaynakça", "Ek", "Kaynakça", len(rows)))
# icindekiler
ws = ICS
yaz(ws, 1, 1, "VitrA E-Ticaret Büyüme Fırsatları · Veri dosyası içindekiler", True); ws.row_dimensions[1].height = 22
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=4)
notlar_yaz(ws, ["Okuma: her sekme rapordaki bir tabloya karşılık gelir; sekme adının başındaki sayı rapordaki bölüm numarası ve bölüm içindeki tablo sırasıdır (ör. 08.2 = Bölüm 08'in ikinci tablosu).",
                "Notlar: her sekmenin A2 hücresinde bölüm, açıklama, kaynak ve sütun açıklamaları yer alır."], 130)
for i, b in enumerate(["Sekme", "Bölüm", "Alt başlık", "Satır"], 1):
    cc = ws.cell(row=4, column=i, value=b); cc.font = F(True, "FFFFFF"); cc.fill = PatternFill("solid", fgColor=BAS); cc.alignment = AL("center"); cc.border = CIZ
for j, (ad, bl, alt, n_) in enumerate(TOC):
    c_ = yaz(ws, 5 + j, 1, ad); c_.hyperlink = "#'%s'!A1" % ad; c_.font = Font(name="Calibri", size=10, color=INK, underline="single")
    yaz(ws, 5 + j, 2, bl); yaz(ws, 5 + j, 3, alt); c4 = yaz(ws, 5 + j, 4, n_, h="center"); c4.number_format = "#,##0"
for i, ww in enumerate([34, 48, 60, 10], 1): ws.column_dimensions[get_column_letter(i)].width = ww
ws.freeze_panes = "A5"
yol = os.path.join(veri.KOK, AD + ".xlsx"); wb.save(yol)
print("kaydedildi:", yol, len(wb.sheetnames), "sekme"); print(wb.sheetnames)
