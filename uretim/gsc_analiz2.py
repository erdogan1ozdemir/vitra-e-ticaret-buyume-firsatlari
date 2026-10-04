# -*- coding: utf-8 -*-
"""Search Console 2. çekimin işlenmesi (04.10.2026) -> veri/islenmis/gsc2.json
- aylık toplam Haz 2025 - Eyl 2026 (cihaz×gün verisinden, tam)
- 12 ay toplamları: 1 Eki 2025 - 30 Eyl 2026
- sayfa türü ve kategori kırılımı: 1 Oca - 30 Eyl 2026 (Aralık 2025'te online.vitra.com.tr adresleri www.vitra.com.tr'ye taşındığı için sayfa düzeyi kırılım 2026'dan başlar)
- blog (/ilham-veren-fikirler/, tüm alt alan adları): aylık seri, 12 ay sayfa listesi, Haz-Eyl yıllık karşılaştırma
- Search Console yapay zeka özellikleri dışa aktarımı (18 May - 29 Eyl 2026, /ilham-veren-fikirler/ filtreli) ile sayfa ve sorgu eşleştirmesi"""
import json, os, re, csv, io, contextlib
from collections import defaultdict
import openpyxl
P = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
H = os.path.join(P, "veri/ham/gsc2")
def J(ad): return json.load(open(os.path.join(H, ad + ".json"), encoding="utf-8"))
with contextlib.redirect_stdout(io.StringIO()):
    import gsc_kategori as GK
import rehber_grup as RG
O = {}
# ------------------------------------------------------------------ aylık toplam
cih = J("gunluk_cihaz")
ay = defaultdict(lambda: [0, 0, 0.0]); gun = defaultdict(lambda: [0, 0]); dv = defaultdict(lambda: [0, 0])
for r in cih:
    d, c = r["keys"]; m = d[:7]
    ay[m][0] += r["clicks"]; ay[m][1] += r["impressions"]; ay[m][2] += r["position"] * r["impressions"]
    gun[d][0] += r["clicks"]; gun[d][1] += r["impressions"]
    if "2025-10-01" <= d <= "2026-09-30": dv[c][0] += r["clicks"]; dv[c][1] += r["impressions"]
O["ay"] = {m: [v[0], v[1], round(v[2] / v[1], 2) if v[1] else None] for m, v in sorted(ay.items())}
O["gun"] = dict(sorted(gun.items()))
Y12 = ["2025-%02d" % m for m in (10, 11, 12)] + ["2026-%02d" % m for m in range(1, 10)]
O["y12"] = {"click": sum(O["ay"][m][0] for m in Y12), "gosterim": sum(O["ay"][m][1] for m in Y12)}
O["cihaz"] = dict(dv)
ul = J("ulke_12ay"); ut = sum(r["clicks"] for r in ul)
O["ulke"] = [(r["keys"][0], r["clicks"], round(100 * r["clicks"] / ut, 1)) for r in ul[:6]]
# ------------------------------------------------------------------ sayfa düzeyi
SA = J("sayfa_aylik")
def tekil(u): return GK.tekil(u)
def yol(u): return re.sub(r"^https?://[^/]+", "", tekil(u)) or "/"
A26 = ["2026-%02d" % m for m in range(1, 10)]
pg = defaultdict(lambda: [0, 0, 0.0])                     # Oca-Eyl 2026 sayfa toplamı
pg12 = defaultdict(lambda: [0, 0, 0.0])                   # 12 ay sayfa toplamı
pgay = defaultdict(lambda: defaultdict(lambda: [0, 0]))   # sayfa × ay
for m, rows in SA.items():
    for r in rows:
        u = tekil(r["keys"][0])
        pgay[u][m][0] += r["clicks"]; pgay[u][m][1] += r["impressions"]
        if m in A26: a = pg[u]; a[0] += r["clicks"]; a[1] += r["impressions"]; a[2] += r["position"] * r["impressions"]
        if m in Y12: a = pg12[u]; a[0] += r["clicks"]; a[1] += r["impressions"]; a[2] += r["position"] * r["impressions"]
TUR = ("kategori", "urun", "koleksiyon", "urun-teknik", "eski-online")
# kategoriye eşlenemeyen sayfalar iki ek gruba ayrılır
def k1_ek(u):
    k = GK.kat(u)
    if k[0] != "Diğer": return k
    p = yol(u)
    if p.startswith(("/v/", "/v-/", "/banyo-koleksiyonlari", "/karo-koleksiyonlari", "/koleksiyon")) or GK.tur(u) == "koleksiyon": return ("Koleksiyon sayfaları", "Koleksiyon")
    return ("Genel ürün listeleri ve kampanyalar", "Genel liste")
def karo_alt(p):
    for rx, ad in ((r"dis-mekan|bahce|teras|havuz|balkon", "Dış mekan karoları"), (r"zemin", "Zemin karoları"), (r"duvar", "Duvar karoları"), (r"mutfak", "Mutfak karoları"), (r"banyo", "Banyo karoları")):
        if re.search(rx, p): return ad
    return "Diğer karo listeleri"
def alt_ad(k, p, t):
    """Kategori içi alt kırılım: alt kategori; karo için kullanım alanı; ek gruplarda sayfa yolu (ürün sayfaları tek kalem)."""
    if k[0] == "Karo Seramik": return "Ürün sayfaları" if t == "urun" else karo_alt(p)
    if k[1] in ("Koleksiyon", "Genel liste"): return "Ürün sayfaları" if t == "urun" else p
    return {"Genel": "Ana kategori ve genel listeler"}.get(k[1], k[1])
tur = defaultdict(lambda: [0, 0, set()]); tur_alt = defaultdict(lambda: defaultdict(int))
kat = defaultdict(lambda: [0, 0, set()]); kat2 = defaultdict(lambda: [0, 0, set()]); kat_alt = defaultdict(lambda: defaultdict(int)); kat_ay = defaultdict(lambda: defaultdict(int))
for u, (c, i, _) in pg.items():
    t = GK.tur(u); a = tur[t]; a[0] += c; a[1] += i; a[2].add(u)
    p = yol(u)
    if t in ("kategori", "urun"): tur_alt[t][GK.kat(u)[0] if GK.kat(u)[0] != "Diğer" else k1_ek(u)[0]] += c
    elif t == "icerik": tur_alt[t][RG.grup_ad(p.rstrip("/").split("/")[-1])[0] if p.count("/") >= 2 else "Dizin sayfası"] += c
    elif t == "servis-bayi": tur_alt[t]["Servisler ve satış noktaları listesi" if p == "/servisler-ve-satis-noktalari" else ("Tek mağaza, bayi ve servis sayfaları" if p.startswith("/servisler-ve-satis-noktalari/") else p)] += c
    else: tur_alt[t][p] += c
    if t in TUR:
        k = k1_ek(u); b = kat[k[0]]; b[0] += c; b[1] += i; b[2].add(u); kat_alt[k[0]][alt_ad(k, p, t)] += c
        b2 = kat2["%s|%s" % k]; b2[0] += c; b2[1] += i; b2[2].add(u)
        for m in A26: kat_ay[k[0]][m] += pgay[u][m][0]
O["tur"] = {t: [v[0], v[1], len(v[2])] for t, v in tur.items()}
O["tur_alt"] = {t: sorted(v.items(), key=lambda x: -x[1])[:6] for t, v in tur_alt.items()}
O["kat"] = {k: [v[0], v[1], len(v[2])] for k, v in kat.items()}
O["kat2"] = {k: [v[0], v[1], len(v[2])] for k, v in kat2.items()}
O["kat_alt"] = {k: sorted(v.items(), key=lambda x: -x[1])[:6] for k, v in kat_alt.items()}
O["kat_ay"] = {k: [v[m] for m in A26] for k, v in kat_ay.items()}
O["top_sayfa"] = [(u, c, i, round(ps / i, 1) if i else None) for u, (c, i, ps) in sorted(pg.items(), key=lambda x: -x[1][0])[:20]]
O["sayfa_toplam_2026"] = sum(v[0] for v in pg.values())
# ------------------------------------------------------------------ blog
def blog_mu(u): return "/ilham-veren-fikirler" in u or "/blog/" in u
AYL = sorted(SA)
O["blog_ay"] = {m: [sum(v[m][0] for u, v in pgay.items() if blog_mu(u)), sum(v[m][1] for u, v in pgay.items() if blog_mu(u))] for m in AYL}
O["site_ay_sayfa"] = {m: sum(v[m][0] for v in pgay.values()) for m in AYL}   # sayfa boyutunda toplam (blog payının paydası)
bl = [(yol(u), c, i, round(ps / i, 1) if i else None) for u, (c, i, ps) in pg12.items() if blog_mu(u)]
O["blog_12ay"] = sorted(bl, key=lambda r: -r[1])
# eski adresler (-old) ayrı sayfa olarak kalır, dizin sayfası ayrıca işaretlenir
# Haz-Eyl karşılaştırması (1 Haz - 29 Eyl)
def sayfa_top(ad):
    d = defaultdict(lambda: [0, 0, 0.0])
    for r in J(ad):
        u = yol(r["keys"][0]); a = d[u]; a[0] += r["clicks"]; a[1] += r["impressions"]; a[2] += r["position"] * r["impressions"]
    return {u: [c, i, round(ps / i, 2) if i else None] for u, (c, i, ps) in d.items()}
B26, B25, BGA = sayfa_top("blog_sayfa_2026"), sayfa_top("blog_sayfa_2025"), sayfa_top("blog_sayfa_genai")
O["blog_hazeyl"] = {"2026": B26, "2025": B25}
# ------------------------------------------------------------------ yapay zeka özellikleri dışa aktarımı
X = "/Users/Erdo/Downloads/vitra.com.tr-Performance-on-Search-Generative-AI-Features-2026-10-04.xlsx"
wb = openpyxl.load_workbook(X, read_only=True)
O["genai_gun"] = {str(d)[:10]: int(v) for d, v in list(wb["Chart"].iter_rows(values_only=True))[1:]}
ga = defaultdict(int)
for u, v in list(wb["Pages"].iter_rows(values_only=True))[1:]: ga[yol(u)] += int(v)
O["genai_sayfa"] = dict(sorted(ga.items(), key=lambda x: -x[1]))
O["genai_cihaz"] = {d: int(v) for d, v in list(wb["Devices"].iter_rows(values_only=True))[1:]}
O["genai_filtre"] = {a: b for a, b in list(wb["Filters"].iter_rows(values_only=True))[1:]}
wb2 = openpyxl.load_workbook(X.replace(".xlsx", " (1).xlsx"), read_only=True)
O["genai_ay"] = [(a, int(b)) for a, b in list(wb2["Chart"].iter_rows(values_only=True))[1:]]
bg = {r["keys"][0]: [r["clicks"], r["impressions"]] for r in J("blog_gunluk_2026")}
O["blog_gun"] = bg
O["blog_gun_2025"] = {r["keys"][0]: [r["clicks"], r["impressions"]] for r in J("blog_gunluk_2025")}
# aylık pay: yapay zeka özelliği gösterimi / blog sayfalarının toplam gösterimi (aynı günler)
pay = []
for etk, top in O["genai_ay"]:
    a, b = etk.split(" - ")
    bi = sum(v[1] for d, v in bg.items() if a <= d <= b); si = sum(v[1] for d, v in gun.items() if a <= d <= b)
    pay.append((a, b, top, bi, si))
O["genai_pay"] = pay
# sayfa düzeyi: yapay zeka özelliği gösterimi, aynı dönemde sayfanın toplam gösterimi ve Haz-Eyl yıllık değişim
sat = []
for u, g in O["genai_sayfa"].items():
    t = BGA.get(u); a6 = B26.get(u); a5 = B25.get(u)
    sat.append({"u": u, "genai": g, "gos_genai_donem": t[1] if t else None, "c26": a6[0] if a6 else 0, "i26": a6[1] if a6 else 0, "p26": a6[2] if a6 else None,
                "c25": a5[0] if a5 else 0, "i25": a5[1] if a5 else 0, "p25": a5[2] if a5 else None})
O["genai_sayfa_tablo"] = sat
# ------------------------------------------------------------------ sorgu düzeyi: Ahrefs SERP özelliği (AI Overview) ile eşleştirme
AH = list(csv.DictReader(open(os.path.join(P, "veri/ham/marka_kelime/ahrefs_blog_aio.tsv"), encoding="utf-8"), delimiter="\t"))
AHD = {r["keyword"].strip().lower(): r for r in AH}
def sorgu_top(ad):
    d = defaultdict(lambda: [0, 0, 0.0])
    for r in J(ad):
        q = r["keys"][0].strip().lower(); a = d[q]; a[0] += r["clicks"]; a[1] += r["impressions"]; a[2] += r["position"] * r["impressions"]
    return d
Q26, Q25 = sorgu_top("blog_sorgu_sayfa_2026"), sorgu_top("blog_sorgu_sayfa_2025")
grp = {"aio": [0, 0, 0.0, 0, 0, 0.0, 0], "yok": [0, 0, 0.0, 0, 0, 0.0, 0]}
qsat = []
for q, r in AHD.items():
    if q not in Q26 and q not in Q25: continue
    a6 = Q26.get(q, [0, 0, 0.0]); a5 = Q25.get(q, [0, 0, 0.0])
    if not (a5[1] and a6[1]): continue   # iki dönemde de gösterim alan sorgular
    g = grp["aio" if "ai_overview" in r["serp_features"].split(",") else "yok"]
    g[0] += a6[0]; g[1] += a6[1]; g[2] += a6[2]; g[3] += a5[0]; g[4] += a5[1]; g[5] += a5[2]; g[6] += 1
    qsat.append({"q": q, "aio": "ai_overview" in r["serp_features"].split(","), "vol": int(r["volume"] or 0), "c26": a6[0], "i26": a6[1], "p26": round(a6[2] / a6[1], 1),
                 "c25": a5[0], "i25": a5[1], "p25": round(a5[2] / a5[1], 1), "url": r["best_position_url"]})
O["aio_grup"] = grp; O["aio_sorgu"] = sorted(qsat, key=lambda r: -(r["c25"]))
O["aio_ahrefs"] = {"kelime": len(AH), "aio": sum(1 for r in AH if "ai_overview" in r["serp_features"].split(",")), "aio_once": sum(1 for r in AH if "ai_overview" in r["serp_features_prev"].split(",")),
                   "min_hacim": min(int(r["volume"]) for r in AH)}
# ------------------------------------------------------------------ GEO bölümü dosyaları (12 ay): rehber sayfaları ve soru sorguları
GEO = os.path.join(P, "veri/ham/geo")
with open(os.path.join(GEO, "gsc_ilham.tsv"), "w", encoding="utf-8") as f:
    for u, c, i, p in O["blog_12ay"]:
        s = u.rstrip("/").split("/")[-1]
        if u.startswith("/ilham-veren-fikirler/") and s != "ilham-veren-fikirler" and c >= 5: f.write("%s\t%d\t%d\t%.1f\n" % (s, c, i, p))
so = {r["keys"][0]: r for r in J("sorgu_12ay")}
SORU = [l.strip() for l in open(os.path.join(P, "veri/kaynak/gsc_soru_liste.txt"), encoding="utf-8") if l.strip()]   # GEO bölümündeki 21 soru sorgusu (sabit liste)
with open(os.path.join(GEO, "gsc_soru.tsv"), "w", encoding="utf-8") as f:
    for q in SORU:
        r = so.get(q)
        if r: f.write("%s\t%d\t%d\t%.1f\n" % (q, r["clicks"], r["impressions"], r["position"]))
        else: print("uyarı · soru sorgusu 12 ay verisinde yok:", q)
json.dump(O, open(os.path.join(P, "veri/islenmis/gsc2.json"), "w", encoding="utf-8"), ensure_ascii=False)
if __name__ == "__main__":
    print("12 ay", O["y12"]); print("ay", {m: v[:2] for m, v in O["ay"].items()})
    print("tür 2026", sorted(O["tur"].items(), key=lambda x: -x[1][0]))
    print("kat 2026", sorted(O["kat"].items(), key=lambda x: -x[1][0]))
    print("blog ay", O["blog_ay"])
    print("genai pay", O["genai_pay"])
    print("aio grup", O["aio_grup"], O["aio_ahrefs"])
